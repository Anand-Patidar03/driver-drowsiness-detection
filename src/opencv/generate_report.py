import docx
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

doc = docx.Document()

# Add Title
title = doc.add_heading('Weekly Progress Report', 0)
title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

doc.add_paragraph('Project: Driver Drowsiness Monitor')
doc.add_paragraph('Role: Computer Vision Module Development')
doc.add_paragraph('Focus: Face & Eye Region Localization (My Individual Contribution)')

# Completed Work
doc.add_heading('Completed Work (Week 1 & 2)', level=1)
p = doc.add_paragraph('During the first two weeks, the focus was primarily on initializing the environment, structuring the codebase, and prototyping the initial face and landmark detection.')

doc.add_heading('Week 1: Environment Setup', level=2)
ul1 = doc.add_paragraph(style='List Bullet')
ul1.add_run('Initialized Python virtual environment to isolate project dependencies.')
ul1 = doc.add_paragraph(style='List Bullet')
ul1.add_run('Configured requirements.txt with essential computer vision packages.')
ul1 = doc.add_paragraph(style='List Bullet')
ul1.add_run('Installed required libraries including OpenCV, NumPy, and MediaPipe/Haar Cascades for facial detection.')

doc.add_heading('Week 2: Face Detection Prototyping', level=2)
ul2 = doc.add_paragraph(style='List Bullet')
ul2.add_run('Created the core src/ directory structure to modularize the work.')
ul2 = doc.add_paragraph(style='List Bullet')
ul2.add_run('Developed initial prototype scripts to verify webcam video stream processing.')
ul2 = doc.add_paragraph(style='List Bullet')
ul2.add_run('Implemented real-time face boundary detection using Haar Cascades to reliably identify the facial region in video frames.')
ul2 = doc.add_paragraph(style='List Bullet')
ul2.add_run('Started setting up the foundation for the reusable eye_detector module.')

# Upcoming Tasks
doc.add_heading('Upcoming and Pending Tasks', level=1)
p2 = doc.add_paragraph('In the coming weeks, the primary goal is to localize the eye regions precisely, crop them from the video stream, and prepare the image tensors for classification while applying unit testing for robustness.')

doc.add_heading('Week 3: Eye Region Localization', level=2)
ul3 = doc.add_paragraph(style='List Bullet')
ul3.add_run('Calculate bounding boxes for the upper half of the face to isolate left and right eye regions.')
ul3 = doc.add_paragraph(style='List Bullet')
ul3.add_run('Map Haar Cascade bounds specifically targeting eyes to improve detection accuracy.')

doc.add_heading('Week 4: Cropping and Formatting', level=2)
ul4 = doc.add_paragraph(style='List Bullet')
ul4.add_run('Implement boundary-safe cropping so the module does not crash if eyes are near the edge of the screen.')
ul4 = doc.add_paragraph(style='List Bullet')
ul4.add_run('Format the cropped eye images into standardized NumPy arrays suitable for direct input into the team’s Machine Learning classifier.')

doc.add_heading('Week 5: Validation & Testing', level=2)
ul5 = doc.add_paragraph(style='List Bullet')
ul5.add_run('Develop unit tests for critical edge cases (e.g., driver looking away, face partially obscured, out-of-bounds crops).')

doc.add_heading('Week 6-7: Polish & Documentation', level=2)
ul6 = doc.add_paragraph(style='List Bullet')
ul6.add_run('Refine the implementation documentation to facilitate seamless integration with the main alert system.')
ul6 = doc.add_paragraph(style='List Bullet')
ul6.add_run('Optimize framerate processing efficiency ahead of the final project handoff.')

doc.save('Progress_Report_Week1_2.docx')
print('Report generated successfully: Progress_Report_Week1_2.docx')
